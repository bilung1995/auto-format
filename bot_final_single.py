import asyncio
import logging
import html
import os
import sqlite3
from datetime import datetime
from pathlib import Path
from calendar import monthrange

from aiogram import Bot, Dispatcher, F
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.filters import Command, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import CallbackQuery, InlineKeyboardButton, InlineKeyboardMarkup, Message

BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()
ADMIN_IDS = {int(x.strip()) for x in os.getenv("ADMIN_IDS", "").split(",") if x.strip().isdigit()}

if not BOT_TOKEN:
    raise RuntimeError("BOT_TOKEN belum diisi di Railway Variables")

DB_PATH = Path("bot.db")

DEFAULT_TEMPLATE = """📍 KAB :
📍 KEC :
📍 KEL :

💰 SALDO :

🚻 KELAMIN :
💳 KPJ :
🎯 SENSOR :
📆 IT :

🏛️ PT :

DPT JMO LASIK ✅"""


def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_database():
    conn = db()
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS users (
        telegram_id INTEGER PRIMARY KEY,
        name TEXT,
        username TEXT,
        balance INTEGER DEFAULT 0,
        created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS subscriptions (
        telegram_id INTEGER PRIMARY KEY,
        package_code TEXT,
        package_name TEXT,
        price INTEGER DEFAULT 0,
        start_date TEXT,
        expiry_date TEXT,
        status TEXT DEFAULT 'inactive'
    );
    CREATE TABLE IF NOT EXISTS transactions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        telegram_id INTEGER NOT NULL,
        amount INTEGER NOT NULL,
        payment_method TEXT DEFAULT 'SEABANK/DANA',
        package_code TEXT,
        package_name TEXT,
        proof_file_id TEXT,
        status TEXT DEFAULT 'pending',
        created_at TEXT NOT NULL,
        processed_at TEXT
    );
    CREATE TABLE IF NOT EXISTS format_settings (
        telegram_id INTEGER PRIMARY KEY,
        template TEXT NOT NULL,
        updated_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS format_results (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        telegram_id INTEGER NOT NULL,
        input_text TEXT,
        result_text TEXT,
        created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS format_history (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        telegram_id INTEGER NOT NULL,
        input_text TEXT,
        result_text TEXT,
        created_at TEXT NOT NULL,
        deleted_at TEXT NOT NULL
    );
    """)
    conn.executescript("""
    CREATE INDEX IF NOT EXISTS idx_transactions_status ON transactions(status);
    CREATE INDEX IF NOT EXISTS idx_transactions_user ON transactions(telegram_id);
    CREATE INDEX IF NOT EXISTS idx_format_results_user ON format_results(telegram_id);
    """)
    conn.commit()
    conn.close()


def register_user(user):
    conn = db()
    conn.execute("""
        INSERT INTO users(telegram_id,name,username,balance,created_at)
        VALUES(?,?,?,0,?)
        ON CONFLICT(telegram_id) DO UPDATE SET
            name=excluded.name, username=excluded.username
    """, (user.id, user.full_name, user.username, datetime.now().isoformat()))
    conn.commit()
    conn.close()


def rupiah(value):
    return f"Rp {int(value):,}".replace(",", ".")


def is_admin(user_id):
    return user_id in ADMIN_IDS


def get_subscription(user_id):
    conn = db()
    row = conn.execute("SELECT * FROM subscriptions WHERE telegram_id=?", (user_id,)).fetchone()
    conn.close()
    return row


def has_auto_format_access(user_id):
    sub = get_subscription(user_id)
    if not sub:
        return False
    if sub["status"] == "unlimited":
        return True
    if sub["status"] == "active" and sub["expiry_date"]:
        try:
            return datetime.now() < datetime.fromisoformat(sub["expiry_date"])
        except ValueError:
            return False
    return False


def add_months(dt, months):
    idx = dt.month - 1 + months
    year = dt.year + idx // 12
    month = idx % 12 + 1
    day = min(dt.day, monthrange(year, month)[1])
    return dt.replace(year=year, month=month, day=day)


def main_menu():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="👤 PROFIL", callback_data="profile"),
         InlineKeyboardButton(text="💳 TOP UP", callback_data="topup")],
        [InlineKeyboardButton(text="📊 STATUS", callback_data="status"),
         InlineKeyboardButton(text="🛠️ SOLUSI JMO", callback_data="solusi_jmo")],
        [InlineKeyboardButton(text="📝 AUTO FORMAT", callback_data="auto_format")],
        [InlineKeyboardButton(text="📞 HUBUNGI ADMIN", callback_data="contact_admin")],
        [InlineKeyboardButton(text="🔐 PANEL ADMIN", callback_data="admin_panel")]
    ])


def back_main():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⬅️ KEMBALI", callback_data="back_main")]
    ])


def auto_menu():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📝 BUAT FORMAT", callback_data="format_create")],
        [InlineKeyboardButton(text="📄 HASIL FORMAT", callback_data="format_results")],
        [InlineKeyboardButton(text="⚙️ SETTING FORMAT", callback_data="format_setting")],
        [InlineKeyboardButton(text="🕘 HISTORY", callback_data="format_history")],
        [InlineKeyboardButton(text="⬅️ KEMBALI", callback_data="back_main")]
    ])


def admin_menu():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="👥 CEK USER AKTIF", callback_data="admin_active")],
        [InlineKeyboardButton(text="💰 LIHAT TRANSAKSI PENDING", callback_data="admin_pending")],
        [InlineKeyboardButton(text="➕ TAMBAH SALDO", callback_data="admin_add_balance"),
         InlineKeyboardButton(text="➖ KURANGI SALDO", callback_data="admin_sub_balance")],
        [InlineKeyboardButton(text="🗑️ HAPUS USER", callback_data="admin_delete_user")],
        [InlineKeyboardButton(text="📢 BROADCAST", callback_data="admin_broadcast")],
        [InlineKeyboardButton(text="⬅️ KEMBALI", callback_data="back_main")]
    ])


class PaymentState(StatesGroup):
    waiting_topup_amount = State()
    waiting_proof = State()


class JmoState(StatesGroup):
    waiting_question = State()


class FormatState(StatesGroup):
    waiting_manual = State()
    waiting_excel = State()
    waiting_setting = State()
    waiting_search = State()


class AdminState(StatesGroup):
    waiting_user_amount = State()
    waiting_delete_user = State()
    waiting_broadcast = State()


dp = Dispatcher()


@dp.message(CommandStart())
async def start(message: Message, state: FSMContext):
    await state.clear()
    register_user(message.from_user)
    await message.answer(
        "🤖 <b>JMO LINTAS TERKINI</b>\n\n"
        f"👋 Selamat datang, <b>{message.from_user.full_name}</b>!\n\n"
        "Silakan pilih menu di bawah:",
        reply_markup=main_menu()
    )


@dp.callback_query(F.data == "back_main")
async def back_main_handler(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    await callback.message.edit_text(
        "🤖 <b>JMO LINTAS TERKINI</b>\n\nSilakan pilih menu:",
        reply_markup=main_menu()
    )
    await callback.answer()


@dp.callback_query(F.data == "profile")
async def profile(callback: CallbackQuery):
    register_user(callback.from_user)
    conn = db()
    user = conn.execute("SELECT * FROM users WHERE telegram_id=?", (callback.from_user.id,)).fetchone()
    sub = conn.execute("SELECT * FROM subscriptions WHERE telegram_id=?", (callback.from_user.id,)).fetchone()
    conn.close()

    username = f"@{user['username']}" if user["username"] else "-"
    package = sub["package_name"] if sub else "Belum ada"
    expiry = "-"
    status = "🔴 Tidak aktif"

    if sub:
        if sub["status"] == "unlimited":
            expiry, status = "SELAMANYA", "🟢 AKTIF"
        elif sub["expiry_date"]:
            exp = datetime.fromisoformat(sub["expiry_date"])
            expiry = exp.strftime("%d-%m-%Y")
            status = "🟢 AKTIF" if datetime.now() < exp else "🔴 EXPIRED"

    await callback.message.edit_text(
        "👤 <b>PROFIL USER</b>\n\n"
        f"🆔 Telegram ID : <code>{user['telegram_id']}</code>\n"
        f"👤 Nama : {user['name']}\n"
        f"📱 Username : {username}\n"
        f"💰 Saldo : <b>{rupiah(user['balance'])}</b>\n\n"
        f"📦 Langganan : {package}\n"
        f"📅 Berakhir : {expiry}\n"
        f"📊 Status : {status}",
        reply_markup=back_main()
    )
    await callback.answer()


@dp.callback_query(F.data == "topup")
async def topup(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    await state.set_state(PaymentState.waiting_topup_amount)
    await callback.message.edit_text(
        "💳 <b>TOP UP SALDO</b>\n\n"
        "Silakan transfer ke:\n\n"
        "🏦 <b>SEABANK</b>\n"
        "901040978290\n"
        "A/N HAMBALI\n\n"
        "💰 <b>DANA</b>\n"
        "083824101264\n"
        "A/N HAMBALI\n\n"
        "Atau bisa ketik nominal.\n"
        "Contoh : <code>50000</code>\n\n"
        "Setelah membayar, ketik nominal yang dibayar lalu tekan <b>SUDAH BAYAR</b>.",
        reply_markup=back_main()
    )
    await callback.answer()


@dp.message(PaymentState.waiting_topup_amount, F.text)
async def topup_amount(message: Message, state: FSMContext):
    raw = message.text.strip().lower().replace("rp", "").replace(".", "").replace(",", "").replace(" ", "")
    if not raw.isdigit() or int(raw) <= 0:
        await message.answer("❌ Nominal tidak valid. Contoh: <code>50000</code>")
        return

    amount = int(raw)
    await state.update_data(amount=amount, package_code=None, package_name=None)
    await state.set_state(PaymentState.waiting_proof)
    await message.answer(
        "✅ <b>NOMINAL TERCATAT</b>\n\n"
        f"💰 Nominal : <b>{rupiah(amount)}</b>\n\n"
        "Setelah transfer, tekan tombol:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="✅ SUDAH BAYAR", callback_data="payment_done")],
            [InlineKeyboardButton(text="❌ BATAL", callback_data="back_main")]
        ])
    )


AUTO_PACKAGES = {
    "af_6m": ("6M", "AUTO FORMAT 6 BULAN", 50000),
    "af_1y": ("1Y", "AUTO FORMAT 1 TAHUN", 80000),
    "af_unlimited": ("UNLIMITED", "AUTO FORMAT UNLIMITED", 200000)
}


@dp.callback_query(F.data == "auto_format")
async def auto_format(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    if has_auto_format_access(callback.from_user.id):
        await callback.message.edit_text(
            "📝 <b>AUTO FORMAT</b>\n\n🔓 Akses kamu aktif.\n\nSilakan pilih menu:",
            reply_markup=auto_menu()
        )
    else:
        await callback.message.edit_text(
            "🔒 <b>AUTO FORMAT TERKUNCI</b>\n\n"
            "Untuk membuka AUTO FORMAT, silakan pilih paket:\n\n"
            "🟢 6 Bulan — <b>Rp50.000</b>\n"
            "🔵 1 Tahun — <b>Rp80.000</b>\n"
            "🟣 Unlimited — <b>Rp200.000</b>",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="🟢 6 BULAN — Rp50.000", callback_data="af_6m")],
                [InlineKeyboardButton(text="🔵 1 TAHUN — Rp80.000", callback_data="af_1y")],
                [InlineKeyboardButton(text="🟣 UNLIMITED — Rp200.000", callback_data="af_unlimited")],
                [InlineKeyboardButton(text="⬅️ KEMBALI", callback_data="back_main")]
            ])
        )
    await callback.answer()


@dp.callback_query(F.data.in_(set(AUTO_PACKAGES.keys())))
async def auto_package(callback: CallbackQuery, state: FSMContext):
    code, name, price = AUTO_PACKAGES[callback.data]
    await state.update_data(amount=price, package_code=code, package_name=name)
    await state.set_state(PaymentState.waiting_proof)

    await callback.message.edit_text(
        "💳 <b>PEMBAYARAN AUTO FORMAT</b>\n\n"
        f"📦 Paket : <b>{name}</b>\n"
        f"💰 Harga : <b>{rupiah(price)}</b>\n\n"
        "Silakan transfer ke:\n\n"
        "🏦 <b>SEABANK</b>\n901040978290\nA/N HAMBALI\n\n"
        "💰 <b>DANA</b>\n083824101264\nA/N HAMBALI\n\n"
        "Setelah membayar, tekan <b>SUDAH BAYAR</b> lalu upload bukti.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="✅ SUDAH BAYAR", callback_data="payment_done")],
            [InlineKeyboardButton(text="⬅️ KEMBALI", callback_data="auto_format")]
        ])
    )
    await callback.answer()


@dp.callback_query(F.data == "payment_done")
async def payment_done(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    if not data.get("amount"):
        await callback.answer("❌ Nominal belum ditentukan.", show_alert=True)
        return
    await state.set_state(PaymentState.waiting_proof)
    await callback.message.edit_text(
        "📸 <b>UPLOAD BUKTI PEMBAYARAN</b>\n\n"
        f"💰 Nominal : <b>{rupiah(data['amount'])}</b>\n"
        f"📦 Paket : {data.get('package_name') or 'TOP UP SALDO'}\n\n"
        "Silakan kirim FOTO bukti transfer di chat ini."
    )
    await callback.answer()


@dp.message(PaymentState.waiting_proof, F.photo)
async def payment_proof(message: Message, state: FSMContext):
    data = await state.get_data()
    amount = int(data.get("amount", 0))
    if amount <= 0:
        await state.clear()
        await message.answer("❌ Nominal pembayaran tidak ditemukan.")
        return

    photo = message.photo[-1]
    conn = db()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO transactions(
            telegram_id,amount,payment_method,package_code,package_name,
            proof_file_id,status,created_at
        ) VALUES(?,?,'SEABANK/DANA',?,?,?,'pending',?)
    """, (
        message.from_user.id, amount, data.get("package_code"),
        data.get("package_name"), photo.file_id, datetime.now().isoformat()
    ))
    tx_id = cur.lastrowid
    conn.commit()
    conn.close()
    await state.clear()

    await message.answer(
        "✅ <b>BUKTI PEMBAYARAN DITERIMA</b>\n\n"
        f"🧾 Transaksi : #{tx_id}\n"
        f"💰 Nominal : {rupiah(amount)}\n"
        f"📦 Paket : {data.get('package_name') or 'TOP UP SALDO'}\n\n"
        "⏳ Menunggu konfirmasi Admin."
    )

    caption = (
        "💰 <b>PEMBAYARAN BARU</b>\n\n"
        f"🧾 Transaksi : #{tx_id}\n"
        f"👤 Nama : {message.from_user.full_name}\n"
        f"🆔 ID : <code>{message.from_user.id}</code>\n"
        f"📱 Username : @{message.from_user.username or '-'}\n"
        f"💰 Nominal : <b>{rupiah(amount)}</b>\n"
        f"📦 Paket : {data.get('package_name') or 'TOP UP SALDO'}\n"
        "🟡 Status : PENDING"
    )
    keyboard = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ SETUJU", callback_data=f"approve_{tx_id}"),
        InlineKeyboardButton(text="❌ TOLAK", callback_data=f"reject_{tx_id}")
    ]])

    for admin_id in ADMIN_IDS:
        try:
            await message.bot.send_photo(admin_id, photo.file_id, caption=caption, reply_markup=keyboard)
        except Exception:
            logging.exception("Gagal mengirim transaksi ke admin")


@dp.message(PaymentState.waiting_proof)
async def payment_wrong_proof(message: Message):
    await message.answer("📸 Silakan kirim FOTO bukti pembayaran.")


async def process_payment(callback: CallbackQuery, tx_id: int, approve: bool):
    if not is_admin(callback.from_user.id):
        await callback.answer("❌ Kamu bukan Admin.", show_alert=True)
        return

    conn = db()
    tx = conn.execute("SELECT * FROM transactions WHERE id=?", (tx_id,)).fetchone()

    if not tx or tx["status"] != "pending":
        conn.close()
        await callback.answer("⚠️ Transaksi sudah diproses/tidak ditemukan.", show_alert=True)
        return

    now = datetime.now().isoformat()

    if approve:
        conn.execute(
            "UPDATE transactions SET status='approved', processed_at=? WHERE id=?",
            (now, tx_id)
        )

        expiry = None
        if tx["package_code"]:
            if tx["package_code"] == "6M":
                expiry, sub_status = add_months(datetime.now(), 6), "active"
            elif tx["package_code"] == "1Y":
                expiry, sub_status = add_months(datetime.now(), 12), "active"
            else:
                sub_status = "unlimited"

            conn.execute("""
                INSERT INTO subscriptions(
                    telegram_id,package_code,package_name,price,start_date,expiry_date,status
                ) VALUES(?,?,?,?,?,?,?)
                ON CONFLICT(telegram_id) DO UPDATE SET
                    package_code=excluded.package_code,
                    package_name=excluded.package_name,
                    price=excluded.price,
                    start_date=excluded.start_date,
                    expiry_date=excluded.expiry_date,
                    status=excluded.status
            """, (
                tx["telegram_id"], tx["package_code"], tx["package_name"], tx["amount"],
                now, expiry.isoformat() if expiry else None, sub_status
            ))
        else:
            conn.execute(
                "UPDATE users SET balance=balance+? WHERE telegram_id=?",
                (tx["amount"], tx["telegram_id"])
            )

        conn.commit()
        conn.close()

        expiry_text = ""
        if tx["package_code"]:
            expiry_text = (
                f"\n📅 Berakhir : "
                f"{expiry.strftime('%d-%m-%Y') if expiry else 'SELAMANYA'}"
            )

        user_text = (
            "✅ <b>PEMBAYARAN DISETUJUI</b>\n\n"
            f"🧾 Transaksi : #{tx_id}\n"
            f"💰 Nominal : {rupiah(tx['amount'])}\n"
            f"📦 Paket : {tx['package_name'] or 'TOP UP SALDO'}"
            f"{expiry_text}\n\n"
            + (
                "🔓 AUTO FORMAT SEKARANG SUDAH TERBUKA."
                if tx["package_code"]
                else "💰 Saldo kamu sudah ditambahkan."
            )
        )
    else:
        conn.execute(
            "UPDATE transactions SET status='rejected', processed_at=? WHERE id=?",
            (now, tx_id)
        )
        conn.commit()
        conn.close()

        user_text = (
            "❌ <b>PEMBAYARAN DITOLAK</b>\n\n"
            f"🧾 Transaksi : #{tx_id}\n"
            f"💰 Nominal : {rupiah(tx['amount'])}\n"
            f"📦 Paket : {tx['package_name'] or 'TOP UP SALDO'}\n\n"
            "Silakan cek kembali bukti pembayaran atau hubungi Admin."
        )

    try:
        await callback.bot.send_message(tx["telegram_id"], user_text)
    except Exception:
        logging.exception("Gagal mengirim notifikasi pembayaran")

    # Pending transactions are normally sent as photos, but support text too.
    try:
        if callback.message.photo:
            old_caption = callback.message.caption or ""
            label = "✅ DISETUJUI" if approve else "❌ DITOLAK"
            await callback.message.edit_caption(
                caption=old_caption + f"\n\n<b>{label}</b>",
                reply_markup=None
            )
        else:
            old_text = callback.message.text or ""
            label = "✅ DISETUJUI" if approve else "❌ DITOLAK"
            await callback.message.edit_text(
                old_text + f"\n\n<b>{label}</b>",
                reply_markup=None
            )
    except Exception:
        logging.exception("Gagal memperbarui pesan admin")

    await callback.answer("✅ Disetujui." if approve else "❌ Ditolak.")


@dp.callback_query(F.data.startswith("approve_"))
async def approve_payment(callback: CallbackQuery):
    try:
        tx_id = int(callback.data.split("_", 1)[1])
    except (ValueError, IndexError):
        await callback.answer("❌ ID transaksi tidak valid.", show_alert=True)
        return
    await process_payment(callback, tx_id, True)


@dp.callback_query(F.data.startswith("reject_"))
async def reject_payment(callback: CallbackQuery):
    try:
        tx_id = int(callback.data.split("_", 1)[1])
    except (ValueError, IndexError):
        await callback.answer("❌ ID transaksi tidak valid.", show_alert=True)
        return
    await process_payment(callback, tx_id, False)


@dp.callback_query(F.data == "status")
async def status(callback: CallbackQuery):
    conn = db()
    user = conn.execute("SELECT * FROM users WHERE telegram_id=?", (callback.from_user.id,)).fetchone()
    sub = conn.execute("SELECT * FROM subscriptions WHERE telegram_id=?", (callback.from_user.id,)).fetchone()
    conn.close()

    balance = user["balance"] if user else 0
    if not sub:
        text = f"📊 <b>STATUS USER</b>\n\n💰 Saldo : <b>{rupiah(balance)}</b>\n📦 Langganan : Belum ada\n📊 Status : 🔴 Tidak aktif"
    elif sub["status"] == "unlimited":
        text = f"📊 <b>STATUS USER</b>\n\n💰 Saldo : <b>{rupiah(balance)}</b>\n📦 Langganan : {sub['package_name']}\n📅 Berakhir : SELAMANYA\n📊 Status : 🟢 AKTIF"
    else:
        exp = datetime.fromisoformat(sub["expiry_date"])
        active = datetime.now() < exp
        text = (
            f"📊 <b>STATUS USER</b>\n\n💰 Saldo : <b>{rupiah(balance)}</b>\n"
            f"📦 Langganan : {sub['package_name']}\n"
            f"📅 Berakhir : {exp.strftime('%d-%m-%Y')}\n"
            f"📊 Status : {'🟢 AKTIF' if active else '🔴 EXPIRED'}"
        )

    await callback.message.edit_text(text, reply_markup=back_main())
    await callback.answer()


# ==================== SOLUSI JMO ====================

def jmo_solution(text):
    q = text.lower()
    rules = [
        (["025", "error 025", "kode 025"],
         "🔎 <b>MASALAH KODE 025</b>\n\n"
         "Periksa data identitas dan proses verifikasi JMO. Pastikan aplikasi terbaru, "
         "koneksi stabil, dan data sesuai dokumen. Jika tetap gagal, coba kembali beberapa saat kemudian."),
        (["login", "tidak bisa login", "gagal login"],
         "🔎 <b>JMO TIDAK BISA LOGIN</b>\n\n"
         "Periksa nomor HP/email dan kata sandi, koneksi internet, lalu update aplikasi JMO."),
        (["verifikasi wajah", "wajah", "face"],
         "🔎 <b>GAGAL VERIFIKASI WAJAH</b>\n\n"
         "Gunakan pencahayaan cukup, wajah terlihat jelas, bersihkan kamera, "
         "dan pastikan izin kamera aktif."),
        (["saldo", "jht tidak muncul", "jht"],
         "🔎 <b>SALDO JHT TIDAK MUNCUL</b>\n\n"
         "Refresh/sinkronkan data, pastikan data kepesertaan sesuai, dan gunakan aplikasi versi terbaru."),
        (["kpj", "tidak ditemukan"],
         "🔎 <b>KPJ TIDAK DITEMUKAN</b>\n\n"
         "Periksa nomor KPJ dan data kepesertaan. Jika memiliki beberapa kepesertaan, periksa semuanya."),
        (["otp", "kode otp"],
         "🔎 <b>OTP BERMASALAH</b>\n\n"
         "Pastikan nomor aktif dan jaringan baik. Tunggu beberapa saat sebelum meminta OTP baru."),
        (["email", "ubah email", "ganti email"],
         "🔎 <b>EMAIL JMO BERMASALAH</b>\n\n"
         "Pastikan email aktif dan dapat menerima pesan. Periksa folder spam dan gunakan email yang benar-benar bisa diakses."),
        (["password", "kata sandi", "lupa sandi", "lupa password"],
         "🔎 <b>LUPA PASSWORD JMO</b>\n\n"
         "Gunakan menu pemulihan/lupa kata sandi di aplikasi JMO. Pastikan nomor atau email yang terdaftar masih aktif."),
        (["aktivasi", "belum terdaftar", "registrasi"],
         "🔎 <b>AKTIVASI/REGISTRASI JMO</b>\n\n"
         "Pastikan data kepesertaan sesuai dan nomor HP/email dapat diakses. Lengkapi proses registrasi sesuai instruksi aplikasi."),
        (["kartu digital", "kartu kepesertaan"],
         "🔎 <b>KARTU KEPESERTAAN DIGITAL</b>\n\n"
         "Pastikan data kepesertaan sudah tersinkron dan aplikasi JMO menggunakan versi terbaru. Coba keluar-masuk kembali jika kartu belum tampil."),
        (["klaim", "pengajuan klaim", "klaim jht"],
         "🔎 <b>MASALAH KLAIM JHT</b>\n\n"
         "Periksa status kepesertaan, kelengkapan data, dan dokumen yang diminta. Jika pengajuan gagal, baca pesan error yang tampil untuk menentukan langkah berikutnya."),
        (["bpu", "bukan penerima upah"],
         "🔎 <b>KEPESERTAAN BPU</b>\n\n"
         "Pastikan data kepesertaan BPU sesuai dan status kepesertaan aktif. Untuk error spesifik, kirim kode errornya."),
        (["data tidak sesuai", "identitas tidak sesuai", "nik"],
         "🔎 <b>DATA IDENTITAS TIDAK SESUAI</b>\n\n"
         "Periksa NIK, nama, tanggal lahir, dan data kepesertaan. Data yang berbeda dapat membuat verifikasi gagal."),
        (["kamera", "izin kamera"],
         "🔎 <b>KAMERA JMO</b>\n\n"
         "Pastikan izin kamera aktif, lensa bersih, pencahayaan cukup, dan wajah terlihat jelas saat proses verifikasi."),
        (["server", "gangguan", "maintenance", "tidak dapat terhubung"],
         "🔎 <b>JMO TIDAK TERHUBUNG</b>\n\n"
         "Periksa koneksi internet dan coba lagi beberapa saat kemudian. Jika gangguan terjadi pada banyak pengguna, kemungkinan layanan sedang bermasalah."),
    ]
    for keys, answer in rules:
        if any(k in q for k in keys):
            return answer
    return (
        "🛠️ <b>SOLUSI JMO</b>\n\n"
        "Saya belum menemukan solusi yang tepat.\n\n"
        "Coba tulis lebih detail, misalnya: <code>Cara atasi 025</code>, "
        "<code>JMO tidak bisa login</code>, atau <code>gagal verifikasi wajah</code>."
    )


@dp.callback_query(F.data == "solusi_jmo")
async def solusi_start(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    await state.set_state(JmoState.waiting_question)
    await callback.message.edit_text(
        "🛠️ <b>SOLUSI JMO</b>\n\n"
        "Silahkan ketikan masalah anda atau upload foto/screenshot.\n\n"
        "Contoh: <code>Cara atasi 025</code>",
        reply_markup=back_main()
    )
    await callback.answer()


@dp.message(JmoState.waiting_question, F.text)
async def solusi_text(message: Message):
    await message.answer(jmo_solution(message.text))


@dp.message(JmoState.waiting_question, F.photo)
async def solusi_photo(message: Message):
    caption = message.caption or ""
    if caption:
        await message.answer(jmo_solution(caption))
    else:
        await message.answer(
            "📸 Foto sudah diterima. Tuliskan kode/error yang terlihat pada foto "
            "(contoh: <code>025</code>) agar solusi bisa dicocokkan."
        )


# ==================== AUTO FORMAT ====================

@dp.callback_query(F.data == "format_create")
async def format_create(callback: CallbackQuery):
    if not has_auto_format_access(callback.from_user.id):
        await callback.answer("🔒 AUTO FORMAT masih terkunci.", show_alert=True)
        return
    await callback.message.edit_text(
        "📝 <b>BUAT FORMAT</b>\n\nPilih cara membuat format:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⌨️ KETIK MANUAL", callback_data="format_manual")],
            [InlineKeyboardButton(text="📊 FILE EXCEL", callback_data="format_excel")],
            [InlineKeyboardButton(text="⬅️ KEMBALI", callback_data="auto_format")]
        ])
    )
    await callback.answer()


@dp.callback_query(F.data == "format_manual")
async def format_manual_start(callback: CallbackQuery, state: FSMContext):
    await state.set_state(FormatState.waiting_manual)
    await callback.message.edit_text(
        "⌨️ <b>FORMAT MANUAL</b>\n\n"
        "Ketik data apa saja. Bot akan mengubahnya ke template SETTING FORMAT.\n\n"
        "Contoh:\n<code>KAB: SERANG\nKEC: CIPOCOK JAYA\nKEL: GELAM\nSALDO: 14 JUTA\nKELAMIN: L\nKPJ: 123\nSENSOR: YA\nIT: 01\nPT: ABC</code>",
        reply_markup=back_main()
    )
    await callback.answer()


def apply_template(template, raw):
    values = {k: "" for k in ["KAB", "KEC", "KEL", "SALDO", "KELAMIN", "KPJ", "SENSOR", "IT", "PT"]}
    lines = [x.strip() for x in raw.splitlines() if x.strip()]

    for key in values:
        for line in lines:
            if ":" in line and line.split(":", 1)[0].strip().upper() == key:
                values[key] = line.split(":", 1)[1].strip()
                break

    result = template
    for key, value in values.items():
        result = result.replace("{" + key + "}", value)

    import re
    for key, value in values.items():
        if value:
            result = re.sub(
                rf"(?im)^([^\n]*{re.escape(key)}\s*:\s*).*$",
                rf"\g<1>{value}",
                result
            )
    return result


@dp.message(FormatState.waiting_manual, F.text)
async def format_manual_receive(message: Message, state: FSMContext):
    data = await state.get_data()
    edit_id = data.get("edit_result_id")

    conn = db()
    if edit_id:
        conn.execute(
            "UPDATE format_results SET result_text=? WHERE id=? AND telegram_id=?",
            (message.text, edit_id, message.from_user.id)
        )
        result = message.text
        result_id = edit_id
    else:
        row = conn.execute(
            "SELECT template FROM format_settings WHERE telegram_id=?",
            (message.from_user.id,)
        ).fetchone()
        template = row["template"] if row else DEFAULT_TEMPLATE
        result = apply_template(template, message.text)
        cur = conn.cursor()
        cur.execute("""
            INSERT INTO format_results(telegram_id,input_text,result_text,created_at)
            VALUES(?,?,?,?)
        """, (message.from_user.id, message.text, result, datetime.now().isoformat()))
        result_id = cur.lastrowid

    conn.commit()
    conn.close()
    await state.clear()

    await message.answer(
        "✅ <b>HASIL FORMAT</b>\n\n<pre>" + html.escape(result[:3900]) + "</pre>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="📋 SALIN", callback_data=f"copy_result_{result_id}"),
             InlineKeyboardButton(text="📋 SALIN SEMUA", callback_data=f"copy_all_{result_id}")],
            [InlineKeyboardButton(text="✏️ EDIT", callback_data=f"edit_result_{result_id}"),
             InlineKeyboardButton(text="🗑️ HAPUS", callback_data=f"delete_result_{result_id}")],
            [InlineKeyboardButton(text="💾 SIMPAN", callback_data=f"save_result_{result_id}")],
            [InlineKeyboardButton(text="⬅️ AUTO FORMAT", callback_data="auto_format")]
        ])
    )


@dp.callback_query(F.data == "format_setting")
async def format_setting(callback: CallbackQuery, state: FSMContext):
    if not has_auto_format_access(callback.from_user.id):
        await callback.answer("🔒 AUTO FORMAT masih terkunci.", show_alert=True)
        return
    conn = db()
    row = conn.execute(
        "SELECT template FROM format_settings WHERE telegram_id=?",
        (callback.from_user.id,)
    ).fetchone()
    conn.close()
    template = row["template"] if row else DEFAULT_TEMPLATE
    await state.set_state(FormatState.waiting_setting)
    await callback.message.edit_text(
        "⚙️ <b>SETTING FORMAT</b>\n\n"
        "Template saat ini:\n\n<pre>" + html.escape(template[:3900]) + "</pre>\n\n"
        "Kirim template baru untuk menggantinya.",
        reply_markup=back_main()
    )
    await callback.answer()


@dp.message(FormatState.waiting_setting, F.text)
async def format_setting_receive(message: Message, state: FSMContext):
    conn = db()
    conn.execute("""
        INSERT INTO format_settings(telegram_id,template,updated_at)
        VALUES(?,?,?)
        ON CONFLICT(telegram_id) DO UPDATE SET
            template=excluded.template, updated_at=excluded.updated_at
    """, (message.from_user.id, message.text, datetime.now().isoformat()))
    conn.commit()
    conn.close()
    await state.clear()
    await message.answer(
        "✅ <b>SETTING FORMAT TERSIMPAN</b>\n\n<pre>" + html.escape(message.text[:3900]) + "</pre>",
        reply_markup=auto_menu()
    )


@dp.callback_query(F.data == "format_excel")
async def format_excel_start(callback: CallbackQuery, state: FSMContext):
    if not has_auto_format_access(callback.from_user.id):
        await callback.answer("🔒 AUTO FORMAT masih terkunci.", show_alert=True)
        return
    await state.set_state(FormatState.waiting_excel)
    await callback.message.edit_text(
        "📊 <b>FORMAT FILE EXCEL</b>\n\nSilakan upload file Excel <b>.xlsx</b>.",
        reply_markup=back_main()
    )
    await callback.answer()


@dp.message(FormatState.waiting_excel, F.document)
async def format_excel_receive(message: Message, state: FSMContext):
    if not (message.document.file_name or "").lower().endswith(".xlsx"):
        await message.answer("❌ File harus berformat .xlsx")
        return
    try:
        from openpyxl import load_workbook
    except ImportError:
        await message.answer("❌ Tambahkan <code>openpyxl</code> ke requirements.txt lalu deploy ulang.")
        return

    file = await message.bot.get_file(message.document.file_id)
    local = Path("/tmp") / f"{message.from_user.id}_{message.document.file_name}"
    await message.bot.download_file(file.file_path, local)

    conn = db()
    row = conn.execute(
        "SELECT template FROM format_settings WHERE telegram_id=?",
        (message.from_user.id,)
    ).fetchone()
    template = row["template"] if row else DEFAULT_TEMPLATE

    wb = load_workbook(local, read_only=True, data_only=True)
    ws = wb.active
    results = []

    for values in ws.iter_rows(values_only=True):
        raw = " ".join(str(v) for v in values if v is not None).strip()
        if raw:
            results.append(apply_template(template, raw))

    if not results:
        conn.close()
        await state.clear()
        await message.answer("❌ File Excel tidak memiliki data.")
        return

    result = "\n\n".join(results)
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO format_results(telegram_id,input_text,result_text,created_at)
        VALUES(?,?,?,?)
    """, (message.from_user.id, message.document.file_name, result, datetime.now().isoformat()))
    rid = cur.lastrowid
    conn.commit()
    conn.close()
    await state.clear()

    await message.answer(
        "✅ <b>HASIL FORMAT EXCEL</b>\n\n<pre>" + result[:3900] + "</pre>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="📋 SALIN SEMUA", callback_data=f"copy_all_{rid}")],
            [InlineKeyboardButton(text="💾 SIMPAN", callback_data=f"save_result_{rid}")],
            [InlineKeyboardButton(text="⬅️ AUTO FORMAT", callback_data="auto_format")]
        ])
    )


@dp.message(FormatState.waiting_excel)
async def format_excel_wrong(message: Message):
    await message.answer("📊 Silakan upload file Excel .xlsx")


@dp.callback_query(F.data == "format_results")
async def format_results(callback: CallbackQuery):
    if not has_auto_format_access(callback.from_user.id):
        await callback.answer("🔒 AUTO FORMAT masih terkunci.", show_alert=True)
        return
    conn = db()
    rows = conn.execute("""
        SELECT id,created_at FROM format_results
        WHERE telegram_id=? ORDER BY id DESC LIMIT 30
    """, (callback.from_user.id,)).fetchall()
    conn.close()

    buttons = [
        [InlineKeyboardButton(
            text=f"📄 #{r['id']} — {r['created_at'][:16].replace('T',' ')}",
            callback_data=f"view_result_{r['id']}"
        )] for r in rows
    ]
    buttons += [
        [InlineKeyboardButton(text="🔎 CARI", callback_data="format_search")],
        [InlineKeyboardButton(text="🗑️ HAPUS HISTORY", callback_data="clear_results")],
        [InlineKeyboardButton(text="⬅️ KEMBALI", callback_data="auto_format")]
    ]

    await callback.message.edit_text(
        "📄 <b>HASIL FORMAT</b>\n\n" + ("Pilih hasil:" if rows else "Belum ada hasil."),
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()


@dp.callback_query(F.data.startswith("view_result_"))
async def view_result(callback: CallbackQuery):
    rid = int(callback.data.split("_")[-1])
    conn = db()
    row = conn.execute(
        "SELECT * FROM format_results WHERE id=? AND telegram_id=?",
        (rid, callback.from_user.id)
    ).fetchone()
    conn.close()

    if not row:
        await callback.answer("Hasil tidak ditemukan.", show_alert=True)
        return

    await callback.message.edit_text(
        "📄 <b>HASIL FORMAT</b>\n\n<pre>" + html.escape(row["result_text"][:3900]) + "</pre>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="📋 SALIN", callback_data=f"copy_result_{rid}"),
             InlineKeyboardButton(text="📋 SALIN SEMUA", callback_data=f"copy_all_{rid}")],
            [InlineKeyboardButton(text="✏️ EDIT", callback_data=f"edit_result_{rid}"),
             InlineKeyboardButton(text="🗑️ HAPUS", callback_data=f"delete_result_{rid}")],
            [InlineKeyboardButton(text="💾 SIMPAN", callback_data=f"save_result_{rid}")],
            [InlineKeyboardButton(text="⬅️ HASIL FORMAT", callback_data="format_results")]
        ])
    )
    await callback.answer()


async def send_copy(callback, text):
    await callback.message.answer("📋 <b>SILAKAN SALIN</b>\n\n<code>" + html.escape(text[:3900]) + "</code>")
    await callback.answer()


@dp.callback_query(F.data.startswith("copy_result_"))
async def copy_result(callback: CallbackQuery):
    rid = int(callback.data.split("_")[-1])
    conn = db()
    row = conn.execute(
        "SELECT result_text FROM format_results WHERE id=? AND telegram_id=?",
        (rid, callback.from_user.id)
    ).fetchone()
    conn.close()
    if row:
        await send_copy(callback, row["result_text"])
    else:
        await callback.answer("Hasil tidak ditemukan.", show_alert=True)


@dp.callback_query(F.data.startswith("copy_all_"))
async def copy_all(callback: CallbackQuery):
    await copy_result(callback)


@dp.callback_query(F.data.startswith("edit_result_"))
async def edit_result(callback: CallbackQuery, state: FSMContext):
    rid = int(callback.data.split("_")[-1])
    conn = db()
    row = conn.execute(
        "SELECT result_text FROM format_results WHERE id=? AND telegram_id=?",
        (rid, callback.from_user.id)
    ).fetchone()
    conn.close()
    if not row:
        await callback.answer("Hasil tidak ditemukan.", show_alert=True)
        return

    await state.update_data(edit_result_id=rid)
    await state.set_state(FormatState.waiting_manual)
    await callback.message.edit_text(
        "✏️ <b>EDIT FORMAT</b>\n\nKirim teks hasil format yang sudah diperbaiki.",
        reply_markup=back_main()
    )
    await callback.answer()


@dp.callback_query(F.data.startswith("delete_result_"))
async def delete_result(callback: CallbackQuery):
    rid = int(callback.data.split("_")[-1])
    conn = db()
    row = conn.execute(
        "SELECT * FROM format_results WHERE id=? AND telegram_id=?",
        (rid, callback.from_user.id)
    ).fetchone()
    if row:
        conn.execute("""
            INSERT INTO format_history(
                telegram_id,input_text,result_text,created_at,deleted_at
            ) VALUES(?,?,?,?,?)
        """, (
            callback.from_user.id, row["input_text"], row["result_text"],
            row["created_at"], datetime.now().isoformat()
        ))
        conn.execute("DELETE FROM format_results WHERE id=? AND telegram_id=?",
                     (rid, callback.from_user.id))
        conn.commit()
    conn.close()
    await callback.message.edit_text(
        "🗑️ <b>HASIL DIHAPUS</b>\n\nHasil masuk ke HISTORY.",
        reply_markup=auto_menu()
    )
    await callback.answer("🗑️ Dihapus.")


@dp.callback_query(F.data.startswith("save_result_"))
async def save_result(callback: CallbackQuery):
    rid = int(callback.data.split("_")[-1])
    conn = db()
    row = conn.execute(
        "SELECT * FROM format_results WHERE id=? AND telegram_id=?",
        (rid, callback.from_user.id)
    ).fetchone()
    if row:
        conn.execute("""
            INSERT INTO format_history(
                telegram_id,input_text,result_text,created_at,deleted_at
            ) VALUES(?,?,?,?,?)
        """, (
            callback.from_user.id, row["input_text"], row["result_text"],
            row["created_at"], datetime.now().isoformat()
        ))
        conn.commit()
    conn.close()
    await callback.answer("💾 Disimpan ke HISTORY.")


@dp.callback_query(F.data == "clear_results")
async def clear_results(callback: CallbackQuery):
    conn = db()
    rows = conn.execute(
        "SELECT * FROM format_results WHERE telegram_id=?",
        (callback.from_user.id,)
    ).fetchall()
    for row in rows:
        conn.execute("""
            INSERT INTO format_history(
                telegram_id,input_text,result_text,created_at,deleted_at
            ) VALUES(?,?,?,?,?)
        """, (
            callback.from_user.id, row["input_text"], row["result_text"],
            row["created_at"], datetime.now().isoformat()
        ))
    conn.execute("DELETE FROM format_results WHERE telegram_id=?", (callback.from_user.id,))
    conn.commit()
    conn.close()
    await callback.message.edit_text(
        "🗑️ <b>SEMUA HASIL FORMAT DIHAPUS</b>\n\nSemua dipindahkan ke HISTORY.",
        reply_markup=auto_menu()
    )
    await callback.answer()


@dp.callback_query(F.data == "format_search")
async def format_search(callback: CallbackQuery, state: FSMContext):
    if not has_auto_format_access(callback.from_user.id):
        await callback.answer("🔒 AUTO FORMAT masih terkunci.", show_alert=True)
        return
    await state.set_state(FormatState.waiting_search)
    await callback.message.edit_text(
        "🔎 <b>CARI HASIL FORMAT</b>\n\n"
        "Ketik kata yang ingin dicari, misalnya:\n"
        "<code>SERANG</code> atau <code>CIPOCOK</code>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⬅️ HASIL FORMAT", callback_data="format_results")]
        ])
    )
    await callback.answer()


@dp.message(FormatState.waiting_search, F.text)
async def format_search_receive(message: Message, state: FSMContext):
    query = message.text.strip()
    if not query:
        await message.answer("❌ Kata pencarian tidak boleh kosong.")
        return

    conn = db()
    rows = conn.execute("""
        SELECT id, result_text, created_at
        FROM format_results
        WHERE telegram_id=? AND result_text LIKE ?
        ORDER BY id DESC LIMIT 20
    """, (message.from_user.id, f"%{query}%")).fetchall()
    conn.close()
    await state.clear()

    if not rows:
        await message.answer(
            f"🔎 Tidak ditemukan hasil dengan kata <b>{query}</b>.",
            reply_markup=auto_menu()
        )
        return

    buttons = [[InlineKeyboardButton(
        text=f"📄 #{r['id']} — {r['created_at'][:16].replace('T',' ')}",
        callback_data=f"view_result_{r['id']}"
    )] for r in rows]
    buttons.append([InlineKeyboardButton(text="⬅️ HASIL FORMAT", callback_data="format_results")])

    await message.answer(
        f"🔎 <b>HASIL PENCARIAN</b>\n\nKata: <code>{query}</code>\n"
        f"Ditemukan: <b>{len(rows)}</b>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )


@dp.callback_query(F.data == "format_history")
async def format_history(callback: CallbackQuery):
    if not has_auto_format_access(callback.from_user.id):
        await callback.answer("🔒 AUTO FORMAT masih terkunci.", show_alert=True)
        return
    conn = db()
    rows = conn.execute("""
        SELECT id,result_text,deleted_at FROM format_history
        WHERE telegram_id=? ORDER BY id DESC LIMIT 30
    """, (callback.from_user.id,)).fetchall()
    conn.close()

    if not rows:
        text = "🕘 <b>HISTORY</b>\n\nBelum ada history."
    else:
        text = "🕘 <b>HISTORY</b>\n\n" + "\n\n".join(
            f"#{r['id']} — {r['deleted_at'][:16].replace('T',' ')}\n<pre>{html.escape(r['result_text'][:500])}</pre>"
            for r in rows
        )

    await callback.message.edit_text(text[:3900], reply_markup=auto_menu())
    await callback.answer()


# ==================== HUBUNGI ADMIN ====================

@dp.callback_query(F.data == "contact_admin")
async def contact_admin(callback: CallbackQuery):
    await callback.message.edit_text(
        "📞 <b>HUBUNGI ADMIN</b>\n\n"
        "Jika membutuhkan bantuan, silakan hubungi Admin:\n\n"
        "👤 <b>Telegram</b>\n"
        "@Hambali1995\n\n"
        "📱 <b>WhatsApp</b>\n"
        "083160776091",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="💬 CHAT TELEGRAM", url="https://t.me/Hambali1995")],
            [InlineKeyboardButton(text="📱 CHAT WHATSAPP", url="https://wa.me/6283160776091")],
            [InlineKeyboardButton(text="⬅️ KEMBALI", callback_data="back_main")]
        ])
    )
    await callback.answer()


# ==================== ADMIN PANEL ====================

@dp.message(Command("admin"))
async def admin_command(message: Message, state: FSMContext):
    if not is_admin(message.from_user.id):
        await message.answer("❌ Kamu bukan Admin.")
        return
    await state.clear()
    await message.answer("🔐 <b>PANEL ADMIN</b>\n\nPilih menu:", reply_markup=admin_menu())


@dp.callback_query(F.data == "admin_cancel")
async def admin_cancel(callback: CallbackQuery, state: FSMContext):
    if not is_admin(callback.from_user.id):
        await callback.answer("❌ Bukan Admin.", show_alert=True)
        return
    await state.clear()
    await callback.message.edit_text(
        "🔐 <b>PANEL ADMIN</b>\n\nPilih menu:",
        reply_markup=admin_menu()
    )
    await callback.answer()


@dp.callback_query(F.data == "admin_panel")
async def admin_panel(callback: CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("❌ Kamu bukan Admin.", show_alert=True)
        return
    await callback.message.edit_text("🔐 <b>PANEL ADMIN</b>\n\nPilih menu:", reply_markup=admin_menu())
    await callback.answer()


@dp.callback_query(F.data == "admin_active")
async def admin_active(callback: CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("❌ Bukan Admin.", show_alert=True)
        return

    conn = db()
    rows = conn.execute("""
        SELECT u.telegram_id,u.name,u.username,s.package_name,s.expiry_date,s.status
        FROM users u JOIN subscriptions s ON s.telegram_id=u.telegram_id
        WHERE s.status='unlimited' OR (s.status='active' AND s.expiry_date>?)
        ORDER BY u.telegram_id
    """, (datetime.now().isoformat(),)).fetchall()
    conn.close()

    if not rows:
        text = "👥 <b>USER AKTIF</b>\n\nTidak ada user aktif."
    else:
        text = "👥 <b>USER AKTIF</b>\n\n" + "\n\n".join(
            f"👤 {r['name']}\n🆔 <code>{r['telegram_id']}</code>\n"
            f"📦 {r['package_name']}\n"
            f"📅 {'SELAMANYA' if r['status']=='unlimited' else r['expiry_date'][:10]}"
            for r in rows
        )
        text += f"\n\n📊 Total user aktif: <b>{len(rows)}</b>"

    await callback.message.edit_text(text[:3900], reply_markup=admin_menu())
    await callback.answer()


@dp.callback_query(F.data == "admin_pending")
async def admin_pending(callback: CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("❌ Bukan Admin.", show_alert=True)
        return

    conn = db()
    rows = conn.execute(
        "SELECT * FROM transactions WHERE status='pending' ORDER BY id DESC"
    ).fetchall()
    conn.close()

    if not rows:
        await callback.message.edit_text(
            "💰 <b>TRANSAKSI PENDING</b>\n\nTidak ada transaksi pending.",
            reply_markup=admin_menu()
        )
        await callback.answer()
        return

    for row in rows:
        keyboard = InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text="✅ SETUJU", callback_data=f"approve_{row['id']}"),
            InlineKeyboardButton(text="❌ TOLAK", callback_data=f"reject_{row['id']}")
        ]])
        caption = (
            "💰 <b>TRANSAKSI PENDING</b>\n\n"
            f"🧾 #{row['id']}\n🆔 <code>{row['telegram_id']}</code>\n"
            f"💰 {rupiah(row['amount'])}\n📦 {row['package_name'] or 'TOP UP SALDO'}"
        )
        if row["proof_file_id"]:
            await callback.bot.send_photo(callback.from_user.id, row["proof_file_id"],
                                          caption=caption, reply_markup=keyboard)
        else:
            await callback.bot.send_message(callback.from_user.id, caption, reply_markup=keyboard)

    await callback.message.edit_text(
        f"💰 <b>TRANSAKSI PENDING</b>\n\nDitemukan <b>{len(rows)}</b> transaksi.",
        reply_markup=admin_menu()
    )
    await callback.answer()


@dp.callback_query(F.data == "admin_add_balance")
async def admin_add_balance_start(callback: CallbackQuery, state: FSMContext):
    if not is_admin(callback.from_user.id):
        await callback.answer("❌ Bukan Admin.", show_alert=True)
        return
    await state.clear()
    await state.set_state(AdminState.waiting_user_amount)
    await state.update_data(action="add")
    await callback.message.edit_text(
        "➕ <b>TAMBAH SALDO</b>\n\nKirim:\n<code>ID_USER NOMINAL</code>\n\nContoh: <code>123456789 50000</code>",
        reply_markup=admin_menu()
    )
    await callback.answer()


@dp.callback_query(F.data == "admin_sub_balance")
async def admin_sub_balance_start(callback: CallbackQuery, state: FSMContext):
    if not is_admin(callback.from_user.id):
        await callback.answer("❌ Bukan Admin.", show_alert=True)
        return
    await state.clear()
    await state.set_state(AdminState.waiting_user_amount)
    await state.update_data(action="sub")
    await callback.message.edit_text(
        "➖ <b>KURANGI SALDO</b>\n\nKirim:\n<code>ID_USER NOMINAL</code>\n\nContoh: <code>123456789 50000</code>",
        reply_markup=admin_menu()
    )
    await callback.answer()


@dp.message(AdminState.waiting_user_amount, F.text)
async def admin_balance_process(message: Message, state: FSMContext):
    if not is_admin(message.from_user.id):
        await state.clear()
        return

    data = await state.get_data()
    parts = message.text.strip().split()
    if len(parts) != 2 or not all(p.isdigit() for p in parts):
        await message.answer("❌ Format salah. Contoh: <code>123456789 50000</code>")
        return

    user_id, amount = int(parts[0]), int(parts[1])
    conn = db()
    user = conn.execute("SELECT * FROM users WHERE telegram_id=?", (user_id,)).fetchone()
    if not user:
        conn.close()
        await message.answer("❌ User tidak ditemukan.")
        return

    if data.get("action") == "add":
        conn.execute("UPDATE users SET balance=balance+? WHERE telegram_id=?", (amount, user_id))
        action_text = f"➕ Ditambah {rupiah(amount)}"
    else:
        conn.execute("UPDATE users SET balance=MAX(balance-?,0) WHERE telegram_id=?", (amount, user_id))
        action_text = f"➖ Dikurangi {rupiah(amount)}"

    new_balance = conn.execute(
        "SELECT balance FROM users WHERE telegram_id=?", (user_id,)
    ).fetchone()["balance"]
    conn.commit()
    conn.close()
    await state.clear()

    await message.answer(
        "✅ <b>BERHASIL</b>\n\n"
        f"👤 User : <code>{user_id}</code>\n{action_text}\n"
        f"💰 Saldo sekarang : <b>{rupiah(new_balance)}</b>",
        reply_markup=admin_menu()
    )
    try:
        await message.bot.send_message(
            user_id,
            "💰 <b>PERUBAHAN SALDO</b>\n\n"
            f"{action_text}\nSaldo sekarang : <b>{rupiah(new_balance)}</b>"
        )
    except Exception:
        pass


@dp.callback_query(F.data == "admin_delete_user")
async def admin_delete_start(callback: CallbackQuery, state: FSMContext):
    if not is_admin(callback.from_user.id):
        await callback.answer("❌ Bukan Admin.", show_alert=True)
        return
    await state.clear()
    await state.set_state(AdminState.waiting_delete_user)
    await callback.message.edit_text(
        "🗑️ <b>HAPUS USER</b>\n\nKirim Telegram ID user.\nContoh: <code>123456789</code>",
        reply_markup=admin_menu()
    )
    await callback.answer()


@dp.message(AdminState.waiting_delete_user, F.text)
async def admin_delete_process(message: Message, state: FSMContext):
    if not is_admin(message.from_user.id):
        await state.clear()
        return
    if not message.text.strip().isdigit():
        await message.answer("❌ Telegram ID harus angka.")
        return

    user_id = int(message.text.strip())
    if user_id in ADMIN_IDS:
        await message.answer("❌ User Admin tidak boleh dihapus.")
        return

    conn = db()
    user = conn.execute("SELECT telegram_id FROM users WHERE telegram_id=?", (user_id,)).fetchone()
    if not user:
        conn.close()
        await message.answer("❌ User tidak ditemukan.")
        return

    for table in ["users", "subscriptions", "transactions", "format_settings", "format_results", "format_history"]:
        conn.execute(f"DELETE FROM {table} WHERE telegram_id=?", (user_id,))
    conn.commit()
    conn.close()
    await state.clear()
    await message.answer(
        f"🗑️ User <code>{user_id}</code> berhasil dihapus.",
        reply_markup=admin_menu()
    )


@dp.callback_query(F.data == "admin_broadcast")
async def admin_broadcast_start(callback: CallbackQuery, state: FSMContext):
    if not is_admin(callback.from_user.id):
        await callback.answer("❌ Bukan Admin.", show_alert=True)
        return
    await state.clear()
    await state.set_state(AdminState.waiting_broadcast)
    await callback.message.edit_text(
        "📢 <b>BROADCAST</b>\n\nKirim pesan/foto/dokumen yang ingin dikirim ke semua user.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="❌ BATAL", callback_data="admin_cancel")]
        ])
    )
    await callback.answer()


@dp.message(AdminState.waiting_broadcast)
async def admin_broadcast_process(message: Message, state: FSMContext):
    if not is_admin(message.from_user.id):
        await state.clear()
        return

    conn = db()
    users = conn.execute("SELECT telegram_id FROM users").fetchall()
    conn.close()

    success = failed = 0
    for row in users:
        try:
            if message.text:
                await message.bot.send_message(row["telegram_id"], message.text)
            elif message.photo:
                await message.bot.send_photo(row["telegram_id"], message.photo[-1].file_id, caption=message.caption or "")
            elif message.document:
                await message.bot.send_document(row["telegram_id"], message.document.file_id, caption=message.caption or "")
            else:
                failed += 1
                continue
            success += 1
            await asyncio.sleep(0.05)
        except Exception:
            failed += 1

    await state.clear()
    await message.answer(
        "📢 <b>BROADCAST SELESAI</b>\n\n"
        f"✅ Berhasil : {success}\n❌ Gagal : {failed}",
        reply_markup=admin_menu()
    )


@dp.message(F.text)
async def fallback(message: Message):
    # Masalah JMO tetap dijawab walaupun user mengetik tanpa membuka menu.
    words = ["jmo", "error", "kode", "025", "login", "verifikasi", "kpj", "jht", "otp", "wajah"]
    if any(w in message.text.lower() for w in words):
        await message.answer(jmo_solution(message.text))
    else:
        await message.answer("Silakan gunakan menu utama dengan /start.", reply_markup=main_menu())


async def main():
    init_database()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
    logging.info("🤖 JMO LINTAS TERKINI sedang berjalan")

    bot = Bot(
        token=BOT_TOKEN,
        default=DefaultBotProperties(parse_mode=ParseMode.HTML)
    )
    try:
        await dp.start_polling(bot)
    finally:
        await bot.session.close()


if __name__ == "__main__":
    asyncio.run(main())
